from io import BytesIO
from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from quarterly_automation.pipeline import apply_reviews, process_sources, validate_sources
from quarterly_automation.workbook import workbook_bytes


def demo_sources():
    return {
        role: pd.read_csv(ROOT / "sample_data" / f"{role}.csv", dtype=str)
        for role in ["current_movements", "asset_master", "projects", "cji3", "vendors", "cost_centers", "annual_movements", "previous_quarter"]
    }


def test_validate_sources_reports_missing_roles_and_columns():
    errors = validate_sources({"current_movements": pd.DataFrame(columns=["AiB"])})

    assert any("current movements" in error for error in errors)
    assert any("missing columns: PSP" in error for error in errors)
    assert any("asset master" in error for error in errors)


def test_pipeline_classifies_and_flags_business_exceptions():
    sources = demo_sources()
    result = process_sources(sources, {"excluded_gl_accounts": ["009999"]}, "Q1 2026")

    assert result.errors == []
    assert len(result.master) == 4
    assert result.summary["excluded_cji3_rows"] == 1
    assert result.master.loc[result.master["AiB"] == "100001", "Classification"].item() == "Intangible Asset"
    assert result.master.loc[result.master["AiB"] == "100002", "Classification"].item() == "Building"
    assert result.master.loc[result.master["AiB"] == "100003", "Used"].item() == "Yes"
    assert result.master.loc[result.master["AiB"] == "100001", "FullYearAmount"].item() == 12000
    assert "Multiple supplier countries" in result.master.loc[result.master["AiB"] == "100001", "ExceptionReason"].item()
    assert "previous quarter" in result.master.loc[result.master["AiB"] == "100001", "ExceptionReason"].item()
    assert result.master.loc[result.master["AiB"] == "100004", "Origin"].item() == "Foreign"


def test_review_decision_updates_master_and_creates_audit_event():
    sources = demo_sources()
    sources["asset_master"].loc[sources["asset_master"]["AiB"] == "100004", "AssetClass"] = "UNKNOWN"
    result = process_sources(sources, {}, "Q1 2026")
    reviewed = result.exceptions.copy()
    reviewed["ReviewerOrigin"] = reviewed["ReviewerOrigin"].replace("", "Domestic")
    target = reviewed["AiB"] == "100004"
    reviewed.loc[target, "ReviewerOrigin"] = "Foreign"
    reviewed.loc[target, "ReviewerClassification"] = "Machine"
    reviewed.loc[target, "ReviewerNotes"] = "Confirmed against project description"

    master, audit = apply_reviews(result.master, reviewed, "Demo reviewer")

    item = master.loc[master["AiB"] == "100004"].iloc[0]
    assert item["Origin"] == "Foreign"
    assert item["Classification"] == "Machine"
    assert item["ExceptionReason"] == ""
    assert len(audit) == len(reviewed)
    assert set(audit["Reviewer"]) == {"Demo reviewer"}


def test_unresolved_review_origin_remains_an_exception():
    sources = demo_sources()
    sources["vendors"] = sources["vendors"].iloc[0:0]
    result = process_sources(sources, {})
    reviewed = result.exceptions.loc[result.exceptions["AiB"] == "100001"].copy()
    reviewed.loc[:, "ReviewerOrigin"] = ""
    reviewed.loc[:, "ReviewerClassification"] = "Machine"

    master, audit = apply_reviews(result.master, reviewed, "Demo reviewer")

    assert (master["Origin"] == "Review").any()
    assert (master["ExceptionReason"] != "").any()
    assert audit.empty


def test_missing_required_source_returns_validation_errors():
    result = process_sources({}, quarter="Q1 2026")

    assert result.master.empty
    assert len(result.errors) == 6


def test_workbook_export_contains_all_expected_sheets():
    result = process_sources(demo_sources(), {}, "Q1 2026")
    output = workbook_bytes(result.master, result.exceptions, result.audit, result.summary)

    workbook = load_workbook(BytesIO(output), read_only=True)

    assert workbook.sheetnames == ["Master_List", "Exception_Review", "Run_Summary", "Audit_Log"]